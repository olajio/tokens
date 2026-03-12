#!/usr/bin/env python3
"""
es_apikey_inventory.py
----------------------
Discovers and inventories all Elasticsearch API keys across all clusters
defined in es_clusters.json. Iterates dev → qa → ccs → prod (or whatever
order the clusters appear in the config file).

Outputs a formatted Excel workbook with:
  - Summary sheet (counts, flags, per-cluster breakdown)
  - All Tokens sheet (full inventory)
  - One sheet per cluster

Usage:
    python es_apikey_inventory.py
    python es_apikey_inventory.py --config /path/to/es_clusters.json
    python es_apikey_inventory.py --output my_inventory.xlsx
    python es_apikey_inventory.py --dump-json
    python es_apikey_inventory.py --list-clusters

Requirements:
    pip install requests openpyxl
"""

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import requests
import urllib3
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

DEFAULT_CONFIG_FILE = "es_clusters.json"


# ──────────────────────────────────────────────
# CONFIG LOADER
# ──────────────────────────────────────────────

def load_config(config_path: str) -> dict:
    path = Path(config_path)
    if not path.exists():
        print(f"❌ Config file not found: {config_path}")
        sys.exit(1)

    with open(path) as f:
        config = json.load(f)

    if "clusters" not in config or not config["clusters"]:
        print("❌ Config must contain a non-empty 'clusters' key.")
        sys.exit(1)

    for name, cluster in config["clusters"].items():
        for field in ("url", "api_key"):
            if field not in cluster:
                print(f"❌ Cluster '{name}' is missing required field: '{field}'")
                sys.exit(1)

    print(f"✅ Loaded config: {config_path}")
    print(f"   Clusters: {', '.join(config['clusters'].keys())}")
    return config


# ──────────────────────────────────────────────
# HELPERS
# ──────────────────────────────────────────────

def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def ms_to_dt(ms: int | None) -> datetime | None:
    if ms is None:
        return None
    return datetime.fromtimestamp(ms / 1000, tz=timezone.utc)


def expiry_status(expiration_ms: int | None, invalidated: bool) -> str:
    if invalidated:
        return "INVALIDATED"
    if expiration_ms is None:
        return "NO EXPIRY ⚠"
    exp = ms_to_dt(expiration_ms)
    if exp < now_utc():
        return "EXPIRED"
    days_left = (exp - now_utc()).days
    if days_left <= 30:
        return f"EXPIRING SOON ({days_left}d)"
    return "ACTIVE"


def summarize_permissions(role_descriptors: dict) -> str:
    if not role_descriptors:
        return "inherited (no inline role)"

    parts = []
    for role_name, role_def in role_descriptors.items():
        indices = role_def.get("indices", [])
        cluster = role_def.get("cluster", [])
        apps    = role_def.get("applications", [])

        role_parts = []
        if cluster:
            role_parts.append(f"cluster:[{','.join(cluster)}]")
        if indices:
            idx = "; ".join(
                f"{','.join(e.get('names', ['*']))} → {','.join(e.get('privileges', []))}"
                for e in indices
            )
            role_parts.append(f"indices:[{idx}]")
        if apps:
            app_str = "; ".join(
                f"{a.get('application','')}:{','.join(a.get('privileges',[]))}"
                for a in apps
            )
            role_parts.append(f"apps:[{app_str}]")

        parts.append(f"{role_name}: " + " | ".join(role_parts) if role_parts else role_name)

    return "\n".join(parts)


# ──────────────────────────────────────────────
# ES API
# ──────────────────────────────────────────────

def fetch_api_keys(cluster_name: str, config: dict) -> list[dict]:
    """Query /_security/api_key for all keys in a cluster."""
    url = f"{config['url'].rstrip('/')}/_security/api_key?with_limited_by=true"
    try:
        resp = requests.get(
            url,
            headers={
                "Authorization": f"ApiKey {config['api_key']}",
                "Content-Type": "application/json",
            },
            verify=config.get("verify_ssl", False),
            timeout=15,
        )
        resp.raise_for_status()
        keys = resp.json().get("api_keys", [])
        print(f"  [{cluster_name}] Found {len(keys)} API key(s)")
        return keys
    except requests.exceptions.ConnectionError:
        print(f"  [{cluster_name}] ❌ Connection failed – is the cluster reachable?")
        return []
    except requests.exceptions.HTTPError as e:
        print(f"  [{cluster_name}] ❌ HTTP error: {e}")
        return []
    except Exception as e:
        print(f"  [{cluster_name}] ❌ Unexpected error: {e}")
        return []


def normalize_row(cluster: str, key: dict) -> dict:
    created_ms  = key.get("creation")
    expiry_ms   = key.get("expiration")
    invalidated = key.get("invalidated", False)

    return {
        "Cluster":                  cluster.upper(),
        "Token Name":               key.get("name", ""),
        "Token ID":                 key.get("id", ""),
        "Creator":                  key.get("username", ""),
        "Realm":                    key.get("realm", ""),
        "Created (UTC)":            ms_to_dt(created_ms).strftime("%Y-%m-%d %H:%M") if created_ms else "",
        "Expires (UTC)":            ms_to_dt(expiry_ms).strftime("%Y-%m-%d %H:%M") if expiry_ms else "Never",
        "Status":                   expiry_status(expiry_ms, invalidated),
        "Permissions":              summarize_permissions(key.get("role_descriptors", {})),
        "Purpose":                  "",
        "Team / Owner":             "",
        "Used In (Script/Process)": "",
        "Source Location":          "",
        "Stored in PMP":            "No",
        "PMP Entry Name":           "",
        "Notes":                    "",
    }


# ──────────────────────────────────────────────
# EXCEL BUILDER
# ──────────────────────────────────────────────

COLUMNS = [
    "Cluster", "Token Name", "Token ID", "Creator", "Realm",
    "Created (UTC)", "Expires (UTC)", "Status", "Permissions",
    "Purpose", "Team / Owner", "Used In (Script/Process)",
    "Source Location", "Stored in PMP", "PMP Entry Name", "Notes",
]

COL_WIDTHS = {
    "Cluster": 10, "Token Name": 30, "Token ID": 26, "Creator": 18,
    "Realm": 14, "Created (UTC)": 18, "Expires (UTC)": 18, "Status": 22,
    "Permissions": 55, "Purpose": 30, "Team / Owner": 20,
    "Used In (Script/Process)": 40, "Source Location": 30,
    "Stored in PMP": 15, "PMP Entry Name": 28, "Notes": 30,
}

STATUS_FILLS = {
    "ACTIVE":        PatternFill("solid", fgColor="C6EFCE"),
    "INVALIDATED":   PatternFill("solid", fgColor="D9D9D9"),
    "EXPIRED":       PatternFill("solid", fgColor="FFC7CE"),
    "NO EXPIRY ⚠":  PatternFill("solid", fgColor="FFEB9C"),
    "EXPIRING SOON": PatternFill("solid", fgColor="FFDDC1"),
}

CLUSTER_HEADER_COLORS = {
    None:   "1F4E79",
    "PROD": "833C00",
    "CCS":  "375623",
    "QA":   "4F3A65",
    "DEV":  "244263",
}


def _status_fill(status: str) -> PatternFill | None:
    for key, fill in STATUS_FILLS.items():
        if status.startswith(key):
            return fill
    return None


def _thin_border() -> Border:
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)


def _build_inventory_sheet(ws, rows: list[dict], cluster_filter: str | None):
    border     = _thin_border()
    hdr_color  = CLUSTER_HEADER_COLORS.get(cluster_filter, "1F4E79")
    hdr_fill   = PatternFill("solid", fgColor=hdr_color)
    manual_fill = PatternFill("solid", fgColor="FFF2CC")
    MANUAL_COLS = {"Purpose", "Team / Owner", "Source Location", "PMP Entry Name"}

    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font      = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        cell.fill      = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = border

    ws.row_dimensions[1].height = 30

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(COLUMNS, start=1):
            value = row.get(col_name, "")
            cell  = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font      = Font(name="Arial", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border    = border

            if col_name == "Status":
                fill = _status_fill(str(value))
                if fill:
                    cell.fill = fill
            elif col_name in MANUAL_COLS and not value:
                cell.fill = manual_fill

        ws.row_dimensions[row_idx].height = 40 if "\n" in str(row.get("Permissions", "")) else 18

    for col_idx, col_name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS.get(col_name, 20)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _build_summary_sheet(ws, rows: list[dict], clusters: list[str]):
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14

    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(bold=True, size=10, name="Arial", color="FFFFFF")
    normal   = Font(size=10, name="Arial")

    ws["A1"] = "Elasticsearch API Key Inventory"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79", name="Arial")
    ws["A2"] = f"Generated: {now_utc().strftime('%Y-%m-%d %H:%M UTC')}"
    ws["A2"].font = Font(italic=True, size=9, name="Arial", color="595959")
    ws["A3"] = "Logstash scan: No"
    ws["A3"].font = Font(italic=True, size=9, name="Arial", color="595959")
    ws.append([])

    ws.append(["Metric", "Count"])
    for cell in ws[ws.max_row]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")

    total         = len(rows)
    active        = sum(1 for r in rows if r["Status"] == "ACTIVE")
    no_expiry     = sum(1 for r in rows if "NO EXPIRY" in r["Status"])
    expired       = sum(1 for r in rows if r["Status"] == "EXPIRED")
    expiring_soon = sum(1 for r in rows if "EXPIRING SOON" in r["Status"])
    invalidated   = sum(1 for r in rows if r["Status"] == "INVALIDATED")
    needs_pmp     = sum(1 for r in rows if r.get("Stored in PMP") == "No" and r["Status"] == "ACTIVE")

    flag_fills = {
        "No Expiry (review!)":  PatternFill("solid", fgColor="FFEB9C"),
        "Expired":              PatternFill("solid", fgColor="FFC7CE"),
        "Expiring Soon (≤30d)": PatternFill("solid", fgColor="FFDDC1"),
        "Not yet in PMP":       PatternFill("solid", fgColor="FFC7CE"),
    }

    for label, count in [
        ("Total API Keys",         total),
        ("Active",                 active),
        ("No Expiry (review!)",    no_expiry),
        ("Expired",                expired),
        ("Expiring Soon (≤30d)",   expiring_soon),
        ("Invalidated",            invalidated),
        ("",                       ""),
        ("Not yet in PMP",         needs_pmp),
    ]:
        ws.append([label, count if count != "" else ""])
        r = ws.max_row
        ws.cell(r, 1).font = normal
        ws.cell(r, 2).font = normal
        ws.cell(r, 2).alignment = Alignment(horizontal="center")
        if label in flag_fills and isinstance(count, int) and count > 0:
            ws.cell(r, 2).fill = flag_fills[label]

    ws.append([])
    ws.append(["Per-Cluster Breakdown", ""])
    for cell in ws[ws.max_row]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")

    for cluster in clusters:
        count = sum(1 for r in rows if r["Cluster"] == cluster)
        ws.append([cluster, count])
        r = ws.max_row
        ws.cell(r, 1).font = Font(bold=True, size=10, name="Arial")
        ws.cell(r, 2).font = normal
        ws.cell(r, 2).alignment = Alignment(horizontal="center")


def build_workbook(rows: list[dict], output_path: str, cluster_names: list[str]):
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    _build_summary_sheet(ws_summary, rows, cluster_names)

    ws_all = wb.create_sheet("All Tokens")
    _build_inventory_sheet(ws_all, rows, cluster_filter=None)

    for cluster in cluster_names:
        cluster_rows = [r for r in rows if r["Cluster"] == cluster]
        if not cluster_rows:
            continue
        ws = wb.create_sheet(cluster)
        _build_inventory_sheet(ws, cluster_rows, cluster_filter=cluster)

    wb.save(output_path)
    print(f"\n✅ Workbook saved → {output_path}")


# ──────────────────────────────────────────────
# MAIN
# ──────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Elasticsearch API Key Inventory (no Logstash scan)"
    )
    parser.add_argument(
        "--config", default=DEFAULT_CONFIG_FILE,
        help=f"Path to cluster config file (default: {DEFAULT_CONFIG_FILE})"
    )
    parser.add_argument(
        "--output", default="es_apikey_inventory.xlsx",
        help="Output Excel filename (default: es_apikey_inventory.xlsx)"
    )
    parser.add_argument(
        "--dump-json", action="store_true",
        help="Also dump raw API responses to es_apikeys_raw.json"
    )
    parser.add_argument(
        "--list-clusters", action="store_true",
        help="Print clusters defined in config and exit"
    )
    args = parser.parse_args()

    config       = load_config(args.config)
    all_clusters = config["clusters"]

    if args.list_clusters:
        print("\nAvailable clusters:")
        print("-" * 50)
        for name, cfg in all_clusters.items():
            print(f"  {name:10} → {cfg['url']}")
            print(f"  {'':10}   {cfg.get('description', 'No description')}\n")
        sys.exit(0)

    print("=" * 60)
    print("  Elasticsearch API Key Inventory")
    print(f"  Config:        {args.config}")
    print(f"  Logstash scan: No")
    print("=" * 60)

    all_rows = {}
    raw_data = {}

    for cluster_name, cluster_cfg in all_clusters.items():
        print(f"\n🔍 Querying [{cluster_name.upper()}] → {cluster_cfg['url']}")
        keys = fetch_api_keys(cluster_name, cluster_cfg)
        raw_data[cluster_name] = keys
        all_rows[cluster_name] = [normalize_row(cluster_name, k) for k in keys]

    flat_rows     = [row for rows in all_rows.values() for row in rows]
    cluster_names = [c.upper() for c in all_clusters.keys()]

    print(f"\n📊 Total keys discovered: {len(flat_rows)}")

    if args.dump_json:
        with open("es_apikeys_raw.json", "w") as f:
            json.dump(raw_data, f, indent=2, default=str)
        print("📄 Raw JSON → es_apikeys_raw.json")

    print("\n📝 Building Excel workbook...")
    build_workbook(flat_rows, args.output, cluster_names)

    no_expiry = [r for r in flat_rows if "NO EXPIRY" in r["Status"]]
    expired   = [r for r in flat_rows if r["Status"] == "EXPIRED"]
    if no_expiry:
        print(f"\n⚠️  {len(no_expiry)} token(s) with NO EXPIRY – review for least-privilege:")
        for r in no_expiry:
            print(f"     [{r['Cluster']}] {r['Token Name']}  (creator: {r['Creator']})")
    if expired:
        print(f"\n🔴 {len(expired)} EXPIRED token(s) – candidates for invalidation:")
        for r in expired:
            print(f"     [{r['Cluster']}] {r['Token Name']}  (creator: {r['Creator']})")


if __name__ == "__main__":
    main()
